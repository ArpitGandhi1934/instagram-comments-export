#!/usr/bin/env python3
"""Export Instagram comments to CSV, Excel (.xlsx) or JSONL.

Calls the yugenox/instagram-comments-scraper Actor on Apify for one or more public posts or reels,
then writes one row per comment: username, text, likes, date, post and whether the comment is by
the post's creator. Filters (likes, keywords, date) run inside the Actor, before billing, so
comments they skip cost nothing.

Run with --help for every option. Requires Python 3.10+ and `pip install -r requirements.txt`
(openpyxl is only needed for .xlsx). APIFY_TOKEN comes from the environment or a .env file next
to this script; --from needs no token.
"""

from __future__ import annotations

import argparse
import csv
import json
import os
import re
import sys
from datetime import datetime
from decimal import Decimal
from pathlib import Path

ACTOR = "yugenox/instagram-comments-scraper"

# Column order for CSV, XLSX and JSONL.
COLUMNS = [
    "username", "text", "likes", "created_at", "post_url", "post_owner", "is_post_owner",
    "is_verified", "comment_id",
]

EPILOG = """\
price: $1.90 per 1,000 comments saved on Apify's Free and Starter plans, $1.50 on Scale,
$1.10 on Business (checked 2026-10-06). Free plans stop at 100 comments per run.

examples:
  python export.py https://www.instagram.com/reel/DdPsDCWRT-u/ --sort top --max-comments 200
  python export.py --file posts.example.txt --keyword price --keyword "how much" --out leads.xlsx
  python export.py --file posts.txt --newer-than "24 hours" --append --out comments.csv
"""


def load_dotenv(path: Path = Path(__file__).with_name(".env")) -> None:
    """Minimal .env reader so the script has no extra dependency."""
    if not path.exists():
        return
    for line in path.read_text().splitlines():
        line = line.strip()
        if line and not line.startswith("#") and "=" in line:
            key, value = line.split("=", 1)
            os.environ.setdefault(key.strip(), value.strip().strip('"').strip("'"))


def _field(obj, camel: str, snake: str):
    """apify-client v1/v2 return dicts, v3 returns model objects. Support both."""
    if obj is None:
        return None
    return obj.get(camel) if isinstance(obj, dict) else getattr(obj, snake, None)


def read_list(path: str) -> list[str]:
    lines = Path(path).read_text(encoding="utf-8").splitlines()
    return [s.strip() for s in lines if s.strip() and not s.strip().startswith("#")]


def is_comment(r: dict) -> bool:
    return "comment_id" in r or r.get("dataType") == "comment" or (
        r.get("dataType") is None and "text" in r and "owner" in r
    )


def to_row(r: dict) -> dict:
    """One Actor row (or one row this script wrote earlier) -> one flat output row."""
    if "comment_id" in r:
        return {c: r.get(c) for c in COLUMNS}
    owner = r.get("owner") or {}
    return {
        "username": owner.get("username"),
        "text": r.get("text"),
        "likes": r.get("likeCount"),
        "created_at": r.get("createdAt"),
        "post_url": r.get("postUrl"),
        "post_owner": r.get("postOwnerUsername"),
        "is_post_owner": r.get("isPostOwner"),
        "is_verified": owner.get("isVerified"),
        "comment_id": str(r["id"]) if r.get("id") is not None else None,
    }


def csv_value(v, col: str):
    """Text that starts with = + - or @ gets a leading ' so Excel and Google Sheets show it as text
    instead of running it as a formula (comments are written by strangers)."""
    if v is None:
        return ""
    if isinstance(v, bool):
        return "true" if v else "false"
    if col in ("text", "username") and isinstance(v, str) and v[:1] in ("=", "+", "-", "@"):
        return "'" + v
    return v


def existing_ids(path: str) -> set[str]:
    p = Path(path)
    if not p.exists() or p.stat().st_size == 0:
        return set()
    if path.lower().endswith(".jsonl"):
        with p.open(encoding="utf-8-sig") as f:
            return {json.loads(line).get("comment_id") for line in f if line.strip()} - {None}
    with p.open(newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        if reader.fieldnames != COLUMNS:
            sys.exit(f"{path} has different columns than this script writes; use a new --out file")
        return {row["comment_id"] for row in reader if row.get("comment_id")}


def write_xlsx(path: str, rows: list[dict]) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
        from openpyxl.styles import Alignment, Font
    except ImportError:
        sys.exit("Writing .xlsx needs openpyxl: pip install openpyxl")

    wb = Workbook()
    ws = wb.active
    ws.title = "Comments"
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in rows:
        created = None
        if r.get("created_at"):
            # Excel has no time zones: the value is UTC.
            created = datetime.fromisoformat(str(r["created_at"]).replace("Z", "+00:00")).replace(tzinfo=None)
        ws.append([None] * len(COLUMNS))
        line = ws.max_row
        values = {
            "username": r.get("username"),
            "text": r.get("text"),
            "likes": r.get("likes"),
            "created_at": created,
            "post_url": r.get("post_url"),
            "post_owner": r.get("post_owner"),
            "is_post_owner": r.get("is_post_owner"),
            "is_verified": r.get("is_verified"),
            "comment_id": r.get("comment_id"),
        }
        for i, col in enumerate(COLUMNS, start=1):
            v = values[col]
            if isinstance(v, str):
                v = ILLEGAL_CHARACTERS_RE.sub("", v)
            cell = ws.cell(row=line, column=i, value=v)
            if isinstance(v, str):
                cell.data_type = "s"  # never a formula, even if the comment starts with =
            if col == "created_at" and v is not None:
                cell.number_format = "yyyy-mm-dd hh:mm"
            elif col == "comment_id":
                cell.number_format = "@"
            elif col == "text":
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    widths = [22, 70, 8, 17, 44, 16, 13, 11, 21]
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(ord("A") + i)].width = w
    ws.freeze_panes = "A2"
    wb.save(path)


def write_out(path: str, rows: list[dict], append: bool) -> None:
    low = path.lower()
    if low.endswith(".xlsx"):
        write_xlsx(path, rows)
        return
    p = Path(path)
    fresh = not append or not p.exists() or p.stat().st_size == 0
    if low.endswith(".jsonl"):
        with p.open("w" if fresh else "a", encoding="utf-8") as f:
            for r in rows:
                f.write(json.dumps(r, ensure_ascii=False, separators=(",", ":")) + "\n")
        return
    # UTF-8 with a BOM, so Excel shows emoji and accents correctly when it opens the file.
    with p.open("w" if fresh else "a", newline="", encoding="utf-8-sig" if fresh else "utf-8") as f:
        w = csv.writer(f)
        if fresh:
            w.writerow(COLUMNS)
        for r in rows:
            w.writerow([csv_value(r.get(c), c) for c in COLUMNS])


def read_from_file(path: str) -> list[dict]:
    text = Path(path).read_text(encoding="utf-8-sig").strip()
    if not text:
        return []
    if text.startswith("["):
        return json.loads(text)
    return [json.loads(line) for line in text.splitlines() if line.strip()]


def print_run_report(client, run) -> None:
    """RUN_REPORT (in the run's key-value store) says what happened to each URL: ok, private,
    not_found, capped... Optional: older runs may not have it."""
    try:
        store_id = _field(run, "defaultKeyValueStoreId", "default_key_value_store_id")
        rec = client.key_value_store(store_id).get_record("RUN_REPORT")
        inputs = ((rec or {}).get("value") or {}).get("queries")
        for it in inputs if isinstance(inputs, list) else []:
            note = f" ({it['note']})" if it.get("note") else ""
            print(f"  {it.get('status', '?')}  {it.get('input', '')}  {it.get('rows', 0)} comments{note}", file=sys.stderr)
    except Exception:  # the report is a convenience; the export doesn't depend on it
        pass


def main() -> None:
    p = argparse.ArgumentParser(
        prog="export.py",
        description="Export Instagram comments to CSV, Excel (.xlsx) or JSONL.",
        epilog=EPILOG,
        formatter_class=argparse.RawDescriptionHelpFormatter,
    )
    p.add_argument("urls", nargs="*", metavar="URL",
                   help="post or reel URLs (/p/<code>/, /reel/<code>/); a profile URL collects comments from its latest posts")
    p.add_argument("--file", metavar="PATH", help="text file with one URL per line (# lines skipped)")
    f = p.add_argument_group("filters (run on Apify before billing; comments they skip are free)")
    f.add_argument("--max-comments", type=int, metavar="N", help="comments to save per post (default 100); with filters, matching comments only")
    f.add_argument("--sort", choices=["newest", "top"], help="newest = Instagram's order (default); top = most-liked first, from up to 2,000 per post")
    f.add_argument("--min-likes", type=int, metavar="N", help="only comments with at least this many likes")
    f.add_argument("--keyword", action="append", default=[], metavar="TEXT", help="only comments containing this word or phrase (case-insensitive); repeat for several")
    f.add_argument("--newer-than", metavar="WHEN", help='only comments after a date (2026-09-01) or within a period ("7 days", "24 hours")')
    o = p.add_argument_group("output")
    o.add_argument("--out", default="comments.csv", metavar="PATH", help="output path: .csv (default comments.csv), .xlsx or .jsonl")
    o.add_argument("--append", action="store_true", help="add only comments whose comment_id is not in --out yet (.csv or .jsonl)")
    o.add_argument("--max-total", type=int, default=1000, metavar="N", help="hard cap on comments across all posts in the run (default 1000)")
    o.add_argument("--max-charge", type=float, default=2.0, metavar="USD", help="hard cost cap for the run in USD, enforced by Apify (default 2)")
    r = p.add_argument_group("re-export without a new run")
    r.add_argument("--dataset", metavar="ID", help="export the dataset of a run you already made (no new run starts; filters are not applied)")
    r.add_argument("--from", dest="from_file", metavar="FILE", help="convert a JSON/JSONL file of the Actor's rows or of this script's rows (no token needed)")
    args = p.parse_args()

    urls = list(args.urls) + (read_list(args.file) if args.file else [])
    if not re.search(r"\.(csv|xlsx|jsonl)$", args.out, re.I):
        p.error("--out must end in .csv, .xlsx or .jsonl")
    if args.append and args.out.lower().endswith(".xlsx"):
        p.error("--append works with .csv or .jsonl (Excel opens the CSV)")
    sources = sum(bool(x) for x in (urls, args.dataset, args.from_file))
    if sources == 0:
        p.error("give post or reel URLs, --file posts.txt, --dataset <id> or --from <file>")
    if sources > 1:
        p.error("use one of: URLs/--file, --dataset, --from")

    if args.from_file:
        items = read_from_file(args.from_file)
    else:
        load_dotenv()
        token = os.environ.get("APIFY_TOKEN")
        if not token:
            sys.exit("Set APIFY_TOKEN (see .env.example), or use --from <file>")
        from apify_client import ApifyClient

        client = ApifyClient(token)
        dataset_id = args.dataset
        if not dataset_id:
            run_input: dict = {"postUrls": urls, "maxItems": args.max_total}
            if args.max_comments is not None:
                run_input["maxComments"] = args.max_comments
            if args.sort:
                run_input["sort"] = args.sort
            if args.min_likes is not None:
                run_input["minLikes"] = args.min_likes
            if args.keyword:
                run_input["keywords"] = args.keyword
            if args.newer_than:
                run_input["newerThan"] = args.newer_than
            print(f"Running {ACTOR} on {len(urls)} URL(s) ...", file=sys.stderr)
            # max_total_charge_usd caps what this run can cost you (pay-per-event Actors only).
            run = client.actor(ACTOR).call(
                run_input=run_input, max_total_charge_usd=Decimal(str(args.max_charge))
            )
            if run is None:
                sys.exit("The run did not finish")
            print(f"Run {_field(run, 'id', 'id')} finished: {_field(run, 'status', 'status')}", file=sys.stderr)
            print_run_report(client, run)
            dataset_id = _field(run, "defaultDatasetId", "default_dataset_id")
        items = list(client.dataset(dataset_id).iterate_items())

    rows = [to_row(it) for it in items if is_comment(it)]
    skipped = 0
    if args.append:
        seen = existing_ids(args.out)
        before = len(rows)
        rows = [r for r in rows if not r.get("comment_id") or r["comment_id"] not in seen]
        skipped = before - len(rows)
    write_out(args.out, rows, args.append)

    verb = "Appended" if args.append else "Wrote"
    extra = f" ({skipped} already there)" if skipped else ""
    print(f"{verb} {len(rows)} comments to {args.out}{extra}", file=sys.stderr)
    asked = min(args.max_total, (args.max_comments or 100) * len(urls))
    if urls and len(items) == 100 and asked > 100:
        print("Exactly 100 comments: Free Apify plans stop at 100 comments per run; a paid plan lifts that.", file=sys.stderr)


if __name__ == "__main__":
    main()
